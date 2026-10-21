"""Bounded XLSX adapter. Never evaluates formulas or follows external links."""
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
import math
import re
import zipfile

from defusedxml import ElementTree as SafeET
from openpyxl import load_workbook

from ..domain.importing import FIELDS, HEADERS, clean

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNPACKED_BYTES = 64 * 1024 * 1024
MAX_ENTRY_BYTES = 16 * 1024 * 1024
MAX_ROWS = 5000
TEMPLATE = "s1-v1"


class ImportFileError(ValueError):
    def __init__(self, code):
        self.code = code


@dataclass(frozen=True)
class ParsedRow:
    number: int
    raw: dict


def inspect_archive(content: bytes, filename: str):
    if not filename.lower().endswith(".xlsx"):
        raise ImportFileError("IMPORT_FILE_TYPE")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("IMPORT_FILE_TOO_LARGE")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [item.filename for item in entries]
            if len(entries) > 1000 or len(set(names)) != len(names):
                raise ImportFileError("IMPORT_ARCHIVE_INVALID")
            if sum(item.file_size for item in entries) > MAX_UNPACKED_BYTES:
                raise ImportFileError("IMPORT_ARCHIVE_LIMIT")
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in entry.filename or entry.flag_bits & 1:
                    raise ImportFileError("IMPORT_ARCHIVE_INVALID")
                if entry.file_size > MAX_ENTRY_BYTES or entry.file_size / max(entry.compress_size, 1) > 250:
                    raise ImportFileError("IMPORT_ARCHIVE_LIMIT")
                if "vbaproject" in entry.filename.lower() or "macrosheet" in entry.filename.lower():
                    raise ImportFileError("IMPORT_MACROS_REJECTED")
                # Reject XML entities before openpyxl parses any shared metadata.
                if entry.filename.endswith((".xml", ".rels")):
                    xml = archive.read(entry)
                    if re.search(br"<!\s*(DOCTYPE|ENTITY)", xml, re.I):
                        raise ImportFileError("IMPORT_ARCHIVE_INVALID")
            if not {"[Content_Types].xml", "xl/workbook.xml", "_rels/.rels"}.issubset(names):
                raise ImportFileError("IMPORT_ARCHIVE_INVALID")
            types = SafeET.fromstring(archive.read("[Content_Types].xml"))
            main = [node.attrib.get("ContentType", "") for node in types if node.attrib.get("PartName") == "/xl/workbook.xml"]
            if main != ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"]:
                raise ImportFileError("IMPORT_FILE_TYPE")
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError("IMPORT_ARCHIVE_INVALID") from None


def scalar(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ImportFileError("IMPORT_CELL_LIMIT")
        if isinstance(value, str) and len(value) > 2000:
            raise ImportFileError("IMPORT_CELL_LIMIT")
        return value
    return str(value)


def read_s1(content: bytes, filename: str, sheet: str, last_row: int):
    inspect_archive(content, filename)
    if not 2 <= last_row <= MAX_ROWS or len(sheet) > 31:
        raise ImportFileError("IMPORT_RANGE_INVALID")
    book = cached = None
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        if sheet not in book.sheetnames:
            raise ImportFileError("IMPORT_SHEET_MISSING")
        values = book[sheet].iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=12)
        saved_values = cached[sheet].iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=12)
        header = next(values)
        next(saved_values)
        if tuple(clean(cell.value) for cell in header) != HEADERS:
            raise ImportFileError("IMPORT_HEADER_MISMATCH")
        rows = []
        for number, (cells, saved) in enumerate(zip(values, saved_values), start=2):
            if not any(cell.value is not None and clean(cell.value) for cell in cells):
                continue
            raw = {}
            for field, cell, cache in zip(FIELDS, cells, saved):
                is_formula = cell.data_type == "f"
                raw[field] = {"value": scalar(cell.value), "formula": scalar(cell.value) if is_formula else None,
                              "cached": scalar(cache.value) if is_formula else None}
            rows.append(ParsedRow(number, raw))
        if not rows:
            raise ImportFileError("IMPORT_NO_ROWS")
        return rows
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError("IMPORT_ARCHIVE_INVALID") from None
    finally:
        if book:
            book.close()
        if cached:
            cached.close()
